"""Write the screened workbook and shortlist CSV."""

import datetime as dt

import pandas as pd
from openpyxl.utils import get_column_letter

from . import config as C

COLUMNS = {
    "rank": "Rank",
    "list": "List",
    "kill_reason": "Kill reason",
    "Auction ID": "Auction ID",
    "Bid4Assets link": "Bid4Assets link",
    "Address (OPA style)": "Address",
    "Unit": "Unit",
    "ZIP": "ZIP",
    "Parcels": "Parcels",
    "property_type": "Property type",
    "building_type": "Building type",
    "livable_sqft": "Livable sq ft",
    "lot_sqft": "Lot sq ft",
    "year_built": "Year built",
    "stories": "Stories",
    "beds": "Beds",
    "zoning": "Zoning",
    "exterior_condition": "OPA exterior condition",
    "condition_band": "Rehab band",
    "violations_open": "Open L&I violations",
    "violations_total": "L&I violations (all)",
    "open_violation_titles": "Open violation types",
    "unsafe": "L&I unsafe",
    "imminently_dangerous": "L&I imminently dangerous",
    "permits_5yr": "Permits (5 yr)",
    "flood_zone": "Flood zone",
    "vacancy_indicator": "City vacancy flag",
    "owner_occupied_hint": "Owner may live there",
    "owner_1": "Owner",
    "last_deed_date": "Last sale date",
    "last_deed_price": "Last sale price",
    "opa_market_value": "OPA market value",
    "as_is_value": "Est. as-is value",
    "arv": "Est. ARV",
    "value_source": "Value source",
    "arv_capped": "ARV capped at 1.25x ZIP p75",
    "comp_confidence": "Comp confidence",
    "comp_count": "Comps used",
    "rehab": "Est. rehab",
    "floor_bid": "Opening bid used",
    "floor_bid_basis": "Opening bid basis",
    "Debt amount": "Debt",
    "max_bid_flip": "Max bid (flip, $50k profit)",
    "max_bid_hold": "Max bid (hold, 12%)",
    "best_play": "Best play",
    "max_bid_to_as_is": "Max bid / as-is",
    "all_in_at_max_bid": "All-in at max bid",
    "uw_best_exit": "Underwriting: best exit",
    "uw_max_bid": "Underwriting: max bid",
    "headroom": "Headroom over opening bid",
    "profit_at_floor": "Flip profit at opening bid",
    "all_in_at_floor": "All-in at opening bid",
    "hold_months": "Hold months",
    "zip_psf_trend": "ZIP $/sf change, 1 yr",
    "score_financial": "Financial",
    "score_physical": "Physical",
    "score_market": "Market",
    "score_upside": "Upside",
    "score_total": "Total score",
    "Sale risk": "Sale risk",
    "Flags": "Intake flags",
    "Plaintiff attorney (normalized)": "Plaintiff attorney",
    "Book/Writ (as exported)": "Book/Writ",
    "OPA (all parcels)": "OPA",
    "Close slot": "Close slot",
}
MONEY = {"Last sale price", "OPA market value", "Est. as-is value", "Est. ARV", "Est. rehab",
         "Opening bid used", "Debt", "Max bid (flip, $50k profit)", "Max bid (hold, 12%)",
         "Headroom over opening bid", "Flip profit at opening bid", "All-in at opening bid",
         "All-in at max bid", "Underwriting: max bid"}


def tidy(m):
    cols = [c for c in COLUMNS if c in m]
    return m[cols].rename(columns=COLUMNS)


def _format(ws, df):
    ws.freeze_panes = "B2"
    for i, col in enumerate(df.columns, 1):
        letter = get_column_letter(i)
        width = min(max(len(str(col)), *(len(str(v)) for v in df[col].head(50))) + 2, 45) if len(df) else len(col) + 2
        ws.column_dimensions[letter].width = width
        if col in MONEY:
            for cell in ws[letter][1:]:
                cell.number_format = '"$"#,##0'
        if col in ("ZIP $/sf change, 1 yr", "Max bid / as-is"):
            for cell in ws[letter][1:]:
                cell.number_format = "0.0%"


def assumptions():
    rows = [(k, v) for k, v in vars(C).items() if k.isupper()]
    return pd.DataFrame(rows, columns=["Setting", "Value"]).astype(str)


def summary(m, source_status):
    counts = m["list"].value_counts() if "list" in m else pd.Series(dtype=int)
    lines = [
        ("Run", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Lots screened", len(m)),
        ("Shortlist", int(counts.get("Shortlist", 0))),
        ("Bench (next in line if shortlist lots get postponed)", int(counts.get("Bench", 0))),
        ("Passed (alive, ranked below bench)", int(counts.get("Pass", 0))),
        ("Killed", int(counts.get("Kill", 0))),
    ]
    lines += [(f"Source: {k}", v) for k, v in source_status.items()]
    lines += [
        ("How to read", "Opening bid used is the floor, not the price you will pay: posted opening bids "
                        "are a median ~10% of the debt. What matters is the max bid, the most you can pay "
                        "and still hit $50k on a flip or 12% on a hold. Max bid / as-is shows how that "
                        "compares with what the house is worth today; the higher, the better your odds "
                        "of winning against other bidders."),
        ("Value estimate", "Where 'Value source' is sold comps: as-is = 25th percentile and ARV = 80th "
                           "percentile $/sf of the nearest 30 open-market sales (Market and comps thread). "
                           "Otherwise: average of OPA market value x ZIP sale-to-assessment ratio and sq ft "
                           "x ZIP median $/sf discounted for condition; ARV = sq ft x ZIP 75th percentile $/sf."),
        ("Hold case", "Rent = 85% of HUD FY2027 Small Area FMR (3BR) for the ZIP, x2 units for multi-family, "
                      "40% expenses and vacancy, 12% yearly cash return on all-in cost. SAFMR is the "
                      "voucher standard and can run above market rent in lower-value ZIPs."),
        ("Not checked here", "Title and lien position, current tax and water balances (the city's "
                             "public tax table is stale), occupancy, interior. Those are the next threads."),
    ]
    return pd.DataFrame(lines, columns=["Item", "Value"])


def write(m, parcels, out_xlsx, out_csv, source_status):
    order = m.sort_values(["rank", "score_total"] if "rank" in m else ["Auction ID"], na_position="last")
    sheets = {
        "Summary": summary(m, source_status),
        "Shortlist": tidy(order[order.get("list") == "Shortlist"]) if "list" in m else pd.DataFrame(),
        "Bench": tidy(order[order.get("list") == "Bench"]) if "list" in m else pd.DataFrame(),
        "Kill list": tidy(order[order.get("list") == "Kill"]) if "list" in m else pd.DataFrame(),
        "All lots": tidy(order),
        "Parcels": parcels,
        "Assumptions": assumptions(),
    }
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            _format(xw.sheets[name], df)
    sheets["Shortlist"].to_csv(out_csv, index=False)
