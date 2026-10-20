"""Build the workbook, recalculate it with LibreOffice, and export the max bids.

Usage:
    python run_model.py [deals.csv] --out-dir DIR

Writes DIR/underwriting_model.xlsx (with calculated values cached, so it
previews correctly) and DIR/max_bids.csv for the bidding sheet.
"""

import argparse
import csv
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import build_model

EXPORT = [
    ("auction_id", "Auction ID"),
    ("address", "Address"),
    ("best_exit", "Best exit"),
    ("max_bid", "Max bid that hits target"),
    ("all_in_cash", "All-in cash at max bid"),
    ("result_at_max", "Profit / return at max bid"),
    ("opening_bid_check", "Check vs opening bid"),
    ("buy_box", "Buy box"),
    ("flip_max_bid", "Flip: max bid"),
    ("hold_max_bid", "Hold: max bid (selected measure)"),
    ("hold_max_bid_yield_on_cost", "Hold: max bid at target yield on cost"),
    ("hold_max_bid_cash_on_cash", "Hold: max bid at target cash-on-cash"),
    ("wholesale_max_bid", "Wholesale: max bid (assign to a flipper, keep the fee)"),
    ("teardown_max_bid", "Teardown: max bid"),
    ("rehab_budget", "Rehab budget (blind)"),
    ("flip_months", "Flip: sale to sold"),
    ("other_exits", "Other exits"),
    ("flags", "Flags"),
]


def recalc(src, out_dir):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(src)],
                       check=True, capture_output=True)
        shutil.move(str(Path(tmp) / Path(src).name), str(out_dir / "underwriting_model.xlsx"))
    return out_dir / "underwriting_model.xlsx"


def export_max_bids(xlsx, out_csv):
    ws = load_workbook(xlsx, data_only=True)["Underwriting"]
    headers = [c.value for c in ws[2]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=build_model.FIRST_ROW, values_only=True):
        if not row[0]:
            continue
        out = {}
        for key, header in EXPORT:
            v = row[idx[header]]
            out[key] = round(v) if isinstance(v, float) and key != "flip_months" else v
        rows.append(out)
    with open(out_csv, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _ in EXPORT])
        w.writeheader()
        w.writerows(rows)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deals", nargs="?")
    ap.add_argument("--out-dir", default=str(build_model.HERE))
    ap.add_argument("--set", action="append", default=[], metavar="NAME=VALUE",
                    help="Override an assumption, e.g. --set FlipMinAnnROI=0.15")
    args = ap.parse_args()
    build_model.set_assumptions(dict(kv.split("=", 1) for kv in args.set))
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    deals = build_model.load_deals(args.deals) if args.deals else build_model.EXAMPLE_DEALS
    with tempfile.TemporaryDirectory() as tmp:
        raw = build_model.build(deals, str(Path(tmp) / "underwriting_model.xlsx"))
        xlsx = recalc(raw, out_dir)
    rows = export_max_bids(xlsx, out_dir / "max_bids.csv")
    print(f"{xlsx}\n{out_dir / 'max_bids.csv'} ({len(rows)} properties)")


if __name__ == "__main__":
    main()
