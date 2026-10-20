"""Write a short business plan per finalist from the recalculated model.

Usage:
    python business_plans.py MODEL_XLSX DEALS_CSV SHORTLIST_CSV -o plans.md [--csv verdicts.csv]
"""

import argparse
import csv

from openpyxl import load_workbook

import build_model

BP = 0.015


def money(v):
    return f"-${-v:,.0f}" if v < 0 else f"${v:,.0f}"


def load_assumptions(wb):
    ws = wb["Assumptions"]
    labels = {item[1]: item[0] for item in build_model.ASSUMPTIONS if item[0]}
    out = {}
    for label, value, *_ in ws.iter_rows(min_row=4, values_only=True):
        if label in labels:
            out[labels[label]] = value
    return out


def verdict(u, deal):
    open_bid = deal["open_bid"] or 0
    mx = u["Max bid that hits target"] or 0
    estimated = "est." in deal["notes"]
    if mx <= 0 or mx < open_bid:
        return "Pass", "max bid is below the opening bid"
    room = mx - open_bid
    over_box = u["Best exit"] == "Flip" and u["All-in cash at max bid"] > 300000
    if room < 15000 or mx < 1.5 * open_bid:
        return "Marginal", f"only {money(room)} of room over the {'estimated ' if estimated else ''}opening bid"
    if over_box:
        return "Bid (over the $300k box)", "all-in cost is above the ideal ceiling, but the profit is large"
    return "Bid", f"{money(room)} of room over the {'estimated ' if estimated else ''}opening bid"


def checks(deal, s, u):
    out = []
    occ = deal["occupancy"]
    if occ == "Occupied":
        out.append("Owner may still live there: look for cars, lights, mail, curtains; plan cash for keys")
    elif occ == "Unknown":
        out.append("Occupancy unknown (modeled as occupied): confirm at the door and with neighbors")
    else:
        out.append("City flags it vacant: check it's sealed and look for squatters or break-ins")
    v = s["Open violation types"] or ""
    if "ROOF" in v or "STRUCT" in v:
        out.append(f"Open L&I violation ({v.lower()}): look at the roofline, parapet, gutters and rear wall")
    else:
        out.append("Roof, front wall, cornice and windows from the street and the rear alley")
    if "CONDO" in s["Building type"]:
        out.append("Condo: get the HOA's monthly dues and 6 months of arrears (they survive); $250/mo is a placeholder")
    if deal["units"] == 2:
        out.append("Count meters and doorbells to confirm two units, and check the rental license and zoning use")
    if s["Comp confidence"] in ("Low", "Medium") or "capped" in deal["notes"]:
        out.append(f"ARV is soft (comps {s['Comp confidence'] or 'n/a'}{', capped' if 'capped' in deal['notes'] else ''}): "
                   "walk the block for renovated sales and list prices")
    if "VACANT LAND" in s["Building type"]:
        out.append("City records call this parcel vacant industrial land: confirm a house actually stands on it")
    if "est." in deal["notes"]:
        out.append("Opening bid not posted yet (21% of debt used): re-check the minimum the morning of Nov 3")
    out.append("Title check: plaintiff in 1st position, tax and water balances, IRS liens, bankruptcy")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("model")
    ap.add_argument("deals")
    ap.add_argument("shortlist")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--csv")
    args = ap.parse_args()

    wb = load_workbook(args.model, data_only=True)
    a = load_assumptions(wb)
    ws = wb["Underwriting"]
    headers = [c.value for c in ws[2]]
    uw = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0]:
            uw[str(row[0])] = dict(zip(headers, row))
    with open(args.deals, newline="") as f:
        deals = {r["auction_id"]: {k: (float(v) if k in ("open_bid", "arv", "rent", "bldg_sf", "units") and v else v)
                                   for k, v in r.items()} for r in csv.DictReader(f)}
    with open(args.shortlist, newline="") as f:
        shortlist = {r["Auction ID"]: r for r in csv.DictReader(f)}

    order = {"Bid": 0, "Bid (over the $300k box)": 1, "Marginal": 2, "Pass": 3}
    rows = []
    for aid, deal in deals.items():
        u = uw[aid]
        v, why = verdict(u, deal)
        rows.append((order[v], -(u["Max bid that hits target"] or 0), aid, v, why))
    rows.sort()

    lines = ["# Business plans: Nov 3, 2026 sheriff sale finalists", "",
             "Prepared Oct 20, 2026 from the underwriting model (`underwriting_model.xlsx`, same folder). "
             "Values are today's sold comps held flat; rents are HUD FY2027 Small Area FMRs; rehab is a blind budget. "
             "Max bid is the most you can pay and still make $50k and 15%/yr on cash on a flip, or 12% yield on cost on a hold. "
             "It's a ceiling, not a target, and it assumes no title surprises.", "",
             "## Summary", "",
             "| # | Address | Verdict | Exit | Opening bid | Max bid | All-in cash | Return at max bid |",
             "|---|---|---|---|---|---|---|---|"]
    bid_cash = 0
    for n, (_, _, aid, v, why) in enumerate(rows, start=1):
        u, d = uw[aid], deals[aid]
        ob = money(d["open_bid"]) + ("*" if "est." in d["notes"] else "")
        lines.append(f"| {n} | {u['Address']} | {v} | {u['Best exit']} | {ob} | {money(u['Max bid that hits target'])} | "
                     f"{money(u['All-in cash at max bid'])} | {u['Profit / return at max bid'] or ''} |")
        if v.startswith("Bid"):
            bid_cash += u["All-in cash at max bid"]
    n_bid = sum(1 for r in rows if r[3].startswith("Bid"))
    lines += ["", "\\* Opening bid not posted yet; 21% of the debt is used as an estimate.", "",
              f"Winning all {n_bid} \"Bid\" lots at their max bids would take about {money(bid_cash)} of cash "
              "over two years, against the $1M cap, so the bidding plan has to pick. "
              "You won't win them all at the max, and many will be postponed.", "",
              "## Plans", ""]

    csv_rows = []
    for n, (_, _, aid, v, why) in enumerate(rows, start=1):
        u, d, s = uw[aid], deals[aid], shortlist[aid]
        occ = int(u["Occupied? (1 = yes)"])
        exit_ = u["Best exit"]
        mx = u["Max bid that hits target"] or 0
        lines.append(f"### {n}. {u['Address']} ({d['zip']}), auction {aid}: {v}")
        lines.append("")
        lines.append(f"{s['Building type'].title()}, {d['bldg_sf']:,.0f} sf, built {d['year_built'] or '?'}, "
                     f"{int(d['units'])} unit{'s' if d['units'] > 1 else ''}. As-is about {money(float(d['asis_value'] or 0))}, "
                     f"ARV {money(d['arv'])}. Verdict: **{v}**, {why}.")
        lines.append("")
        if exit_ == "Pass":
            lines.append("- **Exit:** none hits the targets at any bid above the opening bid.")
        else:
            lines.append(f"- **Best exit:** {exit_}. " + (u["Other exits"] or ""))
        lines.append(f"- **Rehab scope:** {d['scope']}, blind budget {money(u['Rehab budget (blind)'])} "
                     f"(includes 20% contingency and the $10k allowance"
                     + (f", plus {money(float(d['ext_adders']))} for exterior/vacancy" if float(d['ext_adders'] or 0) else "") + ").")
        deed, ej, rm = a["DeedMonths"], a["EjectMonths"] * occ, u["Rehab months"]
        if exit_ == "Value-add hold":
            tl = (f"{deed:.0f} mo to deed" + (f", {ej:.0f} mo to get possession" if ej else "")
                  + f", {rm:.0f} mo rehab, {a['LeaseUpMonths']:.0f} mo to lease: rent starts about month {u['Hold: sale to rent start']:.0f}.")
        else:
            tl = (f"{deed:.0f} mo to deed" + (f", {ej:.0f} mo to get possession" if ej else "")
                  + f", {rm:.0f} mo rehab, {a['SellMonths']:.0f} mo to sell: about {u['Flip: sale to sold']:.0f} months sale to sale.")
        lines.append(f"- **Timeline:** {tl}")
        # Returns at a few bids
        bids = sorted({round(d["open_bid"] or 0, -2), round(mx / 2, -2) if mx else 0, round(mx, -2)})
        cells = []
        for b in bids:
            if b <= 0 and mx <= 0:
                continue
            flip_p = u["Flip: profit before bid"] - b * (1 + BP)
            cash_f = b * (1 + BP) + u["Flip: costs other than bid (pre-sale)"]
            ann = flip_p / cash_f / (u["Flip: sale to sold"] / 12)
            yoc = u["Hold: stabilized NOI / yr"] / (b * (1 + BP) + u["Hold: costs other than bid"])
            cells.append(f"{money(b)} bid: flip {money(flip_p)} ({ann:.0%}/yr), hold {yoc:.1%} yield on cost")
        lines.append("- **Returns:** " + "; ".join(cells) + ".")
        lines.append(f"- **Rent assumed:** {money(d['rent'])}/mo (HUD SAFMR), NOI {money(u['Hold: stabilized NOI / yr'])}/yr.")
        lines.append("- **Check in person and at the courthouse:**")
        for c in checks(d, s, u):
            lines.append(f"  - {c}")
        lines.append("")
        csv_rows.append(dict(auction_id=aid, address=u["Address"], verdict=v, reason=why, best_exit=exit_,
                             opening_bid=d["open_bid"], opening_bid_estimated="est." in d["notes"],
                             max_bid=round(mx), all_in_cash=round(u["All-in cash at max bid"] or 0),
                             result_at_max=u["Profit / return at max bid"] or ""))

    lines += ["## What would change these numbers", "",
              "- **Rehab.** Every plan assumes an interior we haven't seen. Heavy is $85/sf plus 20%. If the drive-by shows a "
              "well-kept exterior, Medium ($55/sf) adds roughly $30k to $50k to the max bid on a 1,200 sf house.",
              "- **Occupancy.** Unknown is modeled as occupied (+5 months, +$5k). Confirmed vacant raises most max bids by $10k to $15k.",
              "- **The 15%/yr floor.** Without it, max bids on the bigger houses rise a lot, but the return on cash drops into single digits.",
              "- **Title.** A 2nd-position plaintiff, unpaid city claims above the bid, or an IRS lien comes straight off the max bid.",
              ""]
    with open(args.out, "w") as f:
        f.write("\n".join(lines))
    if args.csv:
        with open(args.csv, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(csv_rows[0].keys()))
            w.writeheader()
            w.writerows(csv_rows)
    print(args.out)


if __name__ == "__main__":
    main()
